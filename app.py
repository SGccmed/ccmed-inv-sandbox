import os,sqlite3,re,sys,secrets
from datetime import datetime, timezone, date, timedelta
from pathlib import Path
from flask import Flask,request,render_template_string,redirect,url_for,flash,session,g,send_file,abort
from openpyxl import load_workbook
from auth import register_auth,ensure_user_schema,create_user,add_user_block,validate_password,can,roles_of,seed_users_from_env,ADMIN_ROLES
from clients import register_client_routes,ensure_client_schema,resolve_client,serial_matches,serial_key,client_label
import repairs as R  # customer repair tickets: stages, queues, work log (repairs.py)
now_iso=R.now_iso
BASE=Path(__file__).resolve().parent
# The database location can be changed with the CCM_DB_PATH environment variable.
DB=Path(os.environ.get('CCM_DB_PATH') or BASE/'inventory.db')
DB.parent.mkdir(parents=True,exist_ok=True)
def load_secret_key():
 # Use CCM_SECRET if set; otherwise create a random key once and keep it next to the database.
 if os.environ.get('CCM_SECRET'):return os.environ['CCM_SECRET']
 keyfile=DB.parent/'.secret_key'
 if not keyfile.exists():
  keyfile.write_text(secrets.token_hex(32),encoding='utf-8')
  try:os.chmod(keyfile,0o600)
  except OSError:pass
 return keyfile.read_text(encoding='utf-8').strip()
app=Flask(__name__)
app.secret_key=load_secret_key()
app.config.update(
 SESSION_COOKIE_HTTPONLY=True,
 SESSION_COOKIE_SAMESITE='Lax',
 # Secure (HTTPS-only) cookies are on by default; running locally with `python app.py` turns them off.
 SESSION_COOKIE_SECURE=os.environ.get('CCM_SECURE_COOKIES','1')=='1',
 PERMANENT_SESSION_LIFETIME=timedelta(hours=12),
 MAX_CONTENT_LENGTH=100*1024*1024)
# Render and Cloudflare sit in front of the app; trust one proxy hop for the visitor's IP.
from werkzeug.middleware.proxy_fix import ProxyFix
app.wsgi_app=ProxyFix(app.wsgi_app,x_for=1,x_proto=1,x_host=1)
SCHEMA='''CREATE TABLE IF NOT EXISTS units(id INTEGER PRIMARY KEY,source_sheet TEXT,source_row INTEGER,manufacturer TEXT,product_raw TEXT,product_key TEXT,serial TEXT,cost REAL,dom TEXT,notes TEXT,location_raw TEXT,physical_location TEXT,status TEXT,condition TEXT,review TEXT, UNIQUE(source_sheet,source_row)); CREATE TABLE IF NOT EXISTS movements(id INTEGER PRIMARY KEY,unit_id INTEGER,old_location TEXT,new_location TEXT,created_at TEXT DEFAULT CURRENT_TIMESTAMP,note TEXT); CREATE TABLE IF NOT EXISTS repair_orders(id INTEGER PRIMARY KEY,unit_id INTEGER,repair_type TEXT,status TEXT DEFAULT 'Open',notes TEXT,created_at TEXT DEFAULT CURRENT_TIMESTAMP); CREATE TABLE IF NOT EXISTS integration_links(id INTEGER PRIMARY KEY,unit_id INTEGER,system TEXT,external_id TEXT,sync_state TEXT DEFAULT 'not_synced');'''
def conn():
 c=sqlite3.connect(DB,timeout=15);c.row_factory=sqlite3.Row;return c
def norm(s):return re.sub(r'\s+',' ',str(s or '').strip()).upper()
def classify(loc):
 s=norm(loc)
 if s in ('SOLD','S0LD','CLOSED'):return 'Sold archive',None
 if s=='RENTAL':return 'Rental',None
 if s in ('LOAN','LOANER'):return 'On loan',None
 if s in ('PENDING','N'):return 'Pending review',None
 if re.match(r'^[A-Z]\d+(\s*/\s*[A-Z]\d+)?$',s):return 'Active',s.replace(' ','')
 return 'Needs review',None

def import_file(path):
 w=load_workbook(path,read_only=True,data_only=True);s=w['Location'];c=conn();c.executescript(SCHEMA);added=0
 for rownum,row in enumerate(s.iter_rows(min_row=3,values_only=True),start=3):
  loc,prod,serial,cost,dom,notes=(list(row)+[None]*6)[:6]
  if not prod or not str(prod).strip():continue
  status,physical=classify(loc);product=norm(prod);serial=str(serial).strip() if serial is not None else ''
  issues=[]
  if not serial:issues.append('Missing serial')
  if status=='Needs review':issues.append('Unrecognized location/status')
  if status=='Pending review':issues.append('Pending status')
  if cost is None or str(cost).strip()=='':issues.append('Missing cost')
  try:amount=float(cost) if cost is not None and str(cost).strip() else None
  except (TypeError,ValueError):amount=None;issues.append('Invalid cost')
  if amount==0:issues.append('Zero cost (verify)')
  if 'NEW' in product:condition='New (inferred)'
  elif 'REFURB' in product:condition='Refurbished (inferred)'
  else:condition='Unspecified'
  c.execute('INSERT OR IGNORE INTO units(source_sheet,source_row,manufacturer,product_raw,product_key,serial,cost,dom,notes,location_raw,physical_location,status,condition,review) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)',('Location',rownum,None,str(prod).strip(),product,serial,amount,str(dom or ''),str(notes or ''),str(loc or ''),physical,status,condition,'; '.join(issues)))
  added+=c.execute('SELECT changes()').fetchone()[0]
 # Flag duplicates for review, not merge
 duplicates=c.execute("SELECT UPPER(TRIM(serial)) serial,COUNT(*) n FROM units WHERE TRIM(serial)<>'' GROUP BY UPPER(TRIM(serial)) HAVING COUNT(*)>1").fetchall()
 for d in duplicates:
  c.execute("UPDATE units SET review=TRIM(COALESCE(review,'')||'; Duplicate serial (verify)', '; ') WHERE UPPER(TRIM(serial))=?",(d['serial'],))
 c.commit();c.close();return added,len(duplicates)
# New inventory entry: a single shared, never-reset sequence across CC, RP and CR.
ENTRY_SCHEMA="""CREATE TABLE IF NOT EXISTS inventory_id_sequence (
 seq INTEGER PRIMARY KEY AUTOINCREMENT, prefix TEXT NOT NULL,
 generated_id TEXT UNIQUE, created_at TEXT NOT NULL);
 CREATE TABLE IF NOT EXISTS core_items (
 id INTEGER PRIMARY KEY, inventory_code TEXT NOT NULL UNIQUE,
 manufacturer TEXT, product_raw TEXT NOT NULL, serial TEXT, location TEXT,
 supplier TEXT, original_cost REAL NOT NULL DEFAULT 0,
 stage TEXT NOT NULL DEFAULT 'Needs Repair', notes TEXT,
 qc_result TEXT, qc_technician TEXT, qc_date TEXT, qc_notes TEXT,
 converted_unit_id INTEGER, created_at TEXT NOT NULL);
 CREATE TABLE IF NOT EXISTS core_transactions (
 id INTEGER PRIMARY KEY, core_id INTEGER NOT NULL,
 performed_at TEXT NOT NULL, technician TEXT, description TEXT NOT NULL,
 labor_method TEXT NOT NULL DEFAULT 'manual', labor_hours REAL NOT NULL DEFAULT 0,
 labor_rate REAL NOT NULL DEFAULT 0, labor_cost REAL NOT NULL DEFAULT 0,
 parts_cost REAL NOT NULL DEFAULT 0, other_cost REAL NOT NULL DEFAULT 0);
 CREATE TABLE IF NOT EXISTS receiving_batches (
 id INTEGER PRIMARY KEY, created_at TEXT NOT NULL, item_count INTEGER NOT NULL);
"""

def ensure_entry_schema(c):
 c.executescript(ENTRY_SCHEMA)
 cols={x['name'] for x in c.execute('PRAGMA table_info(units)')}
 for name,typ in [('inventory_code','TEXT'),('supplier','TEXT'),('received_date','TEXT'),('entry_batch_id','INTEGER')]:
  if name not in cols:c.execute('ALTER TABLE units ADD COLUMN '+name+' '+typ)
 cols={x['name'] for x in c.execute('PRAGMA table_info(customer_repairs)')}
 if 'inventory_code' not in cols:c.execute('ALTER TABLE customer_repairs ADD COLUMN inventory_code TEXT')
 # Existing RMA records are retained; IDs are assigned only to new entries.
 c.execute('CREATE UNIQUE INDEX IF NOT EXISTS idx_unit_inventory_code ON units(inventory_code)')
 c.execute('CREATE UNIQUE INDEX IF NOT EXISTS idx_repair_inventory_code ON customer_repairs(inventory_code)')

def new_inventory_code(c,prefix,received_date=None):
 if prefix not in ('CC','RP','CR'):raise ValueError('Unknown inventory prefix')
 today=date.today().strftime('%m%d%y') # Entry date, not historical received date
 now=now_iso()
 cur=c.execute('INSERT INTO inventory_id_sequence(prefix,created_at) VALUES (?,?)',(prefix,now))
 seq=cur.lastrowid
 code=f'{prefix}-{today}-{seq:03d}'
 c.execute('UPDATE inventory_id_sequence SET generated_id=? WHERE seq=?',(code,seq))
 return code

def amount(value,label,required=False):
 raw=str(value or '').strip()
 if not raw:return 0.0 if not required else None
 try:
  n=float(raw)
  if not (0<=n<1e12):raise ValueError()
  return round(n,2)
 except ValueError:raise ValueError(label+' must be a nonnegative number')

def entry_row(f):
 return {k:str(f.get(k,'') or '').strip() for k in ('type','manufacturer','product','serial','location','supplier','cost','condition','notes','customer','rma','technician','received_date')}

@app.route('/inventory/new',methods=['GET','POST'])
def inventory_new():
 if request.method=='GET':return render_template_string(ENTRY_TEMPLATE,today=date.today().isoformat())
 # Multiple-row entry is atomic: invalid rows do not result in partial receiving.
 f=request.form;fields=('type','manufacturer','product','serial','location','supplier','cost','condition','notes','customer','rma','technician','received_date')
 length=max((len(f.getlist(k+'[]')) for k in fields),default=0)
 rows=[];c=conn();ensure_entry_schema(c)
 for i in range(length):
  row={k:(f.getlist(k+'[]')[i].strip() if i<len(f.getlist(k+'[]')) else '') for k in fields}
  if not any(row.values()):continue
  n=len(rows)+1
  if not row['type'] or row['type'] not in ('CC','CR','RP'):c.close();return f'Row {n}: select CC, CR or RP',400
  if not row['product']:c.close();return f'Row {n}: model / product is required',400
  if row['type']=='RP' and not row['customer']:c.close();return f'Row {n}: a client is required for RP items (the RMA number is created automatically if left blank)',400
  row['client']=None
  if row['type']=='RP':
   row['client']=resolve_client(c,row['customer'])
   if row['client'] is None:c.close();return f'Row {n}: customer "{row["customer"]}" is not in the client list. Pick one from the list or add it under Clients first.',400
  try:
   if row['received_date']:date.fromisoformat(row['received_date'])
   row['cost_value']=amount(row['cost'],'Purchase cost')
  except ValueError as e:c.close();return f'Row {n}: '+str(e),400
  rows.append(row)
 if not rows:c.close();return 'Enter at least one item',400
 # Serial numbers are not unique keys (different manufacturers can repeat them), so matches only raise an alert.
 alerts=[];seen={}
 for i,row in enumerate(rows,1):
  k=serial_key(row['serial'])
  if not k:continue
  for m in serial_matches(c,row['serial']):alerts.append(f"S/N {row['serial']} (row {i}) also on {m['code']} — {m['manufacturer']} {m['product']} [{m['status']}]")
  if k in seen:alerts.append(f"S/N {row['serial']} entered on rows {seen[k]} and {i}")
  seen.setdefault(k,i)
 try:
  c.execute('BEGIN IMMEDIATE')
  batch=c.execute('INSERT INTO receiving_batches(created_at,item_count) VALUES (?,?)',(now_iso(),len(rows))).lastrowid
  codes=[]
  for row in rows:
   typ=row['type'];code=new_inventory_code(c,typ);codes.append(code)
   received=row['received_date'] or date.today().isoformat()
   if typ=='CC':
    c.execute("""INSERT INTO units(manufacturer,product_raw,product_key,serial,cost,dom,notes,location_raw,physical_location,status,condition,review,inventory_code,supplier,received_date,entry_batch_id)
     VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
     (row['manufacturer'],row['product'],norm(row['product']),row['serial'],row['cost_value'],'',row['notes'],row['location'],row['location'],'Active',row['condition'] or 'Unspecified','',code,row['supplier'],received,batch))
   elif typ=='CR':
    c.execute("""INSERT INTO core_items(inventory_code,manufacturer,product_raw,serial,location,supplier,original_cost,notes,created_at)
     VALUES (?,?,?,?,?,?,?,?,?)""",(code,row['manufacturer'],row['product'],row['serial'],row['location'],row['supplier'],row['cost_value'],row['notes'],now_iso()))
   else:
    # Customer repair ticket (see repairs.py); a blank RMA gets the next RMA-##### number.
    R.create_ticket(c,client=row['client'],rma=row['rma'],received=received,code=code,
     fields={'manufacturer':row['manufacturer'],'model':row['product'],'serial':row['serial'],'location':row['location'],'notes':row['notes']},
     assigned_to=row['technician'],note='Received through New Inventory')
  c.commit()
 except (sqlite3.Error,ValueError) as e:
  c.rollback();c.close()
  if 'rma' in str(e).lower():return 'Receiving not saved: an RMA number on this form is already used on another ticket.',400
  return 'Receiving not saved: '+str(e),400
 c.close();flash('Saved '+str(len(codes))+' items: '+', '.join(codes))
 for a in alerts:flash('⚠ Serial number alert: '+a)
 return redirect(url_for('inventory_new'))

@app.route('/cores')
def cores():
 c=conn();ensure_entry_schema(c)
 rows=c.execute("""SELECT c.*,COALESCE((SELECT SUM(labor_cost+parts_cost+other_cost) FROM core_transactions WHERE core_id=c.id),0) repair_cost FROM core_items c ORDER BY c.id DESC""").fetchall();c.close()
 return render_template_string(CORE_TEMPLATE,rows=rows,detail=None,transactions=[],total=0)

@app.route('/cores/<int:cid>',methods=['GET','POST'])
def core_detail(cid):
 c=conn();ensure_entry_schema(c);core=c.execute('SELECT * FROM core_items WHERE id=?',(cid,)).fetchone()
 if core is None:c.close();return 'Core not found',404
 if request.method=='POST':
  f=request.form;action=f.get('action','')
  try:
   c.execute('BEGIN IMMEDIATE')
   if core['stage']=='Converted to Stock':raise ValueError('Converted cores cannot be modified')
   if action=='transaction':
    method=f.get('labor_method','manual');hours=amount(f.get('hours'),'Hours');rate=amount(f.get('rate'),'Hourly rate')
    labor=amount(f.get('labor'),'Labor cost') if method=='manual' else round(hours*rate,2)
    parts=amount(f.get('parts'),'Parts cost');other=amount(f.get('other'),'Other cost')
    if not f.get('description','').strip():raise ValueError('Transaction description required')
    c.execute("""INSERT INTO core_transactions(core_id,performed_at,technician,description,labor_method,labor_hours,labor_rate,labor_cost,parts_cost,other_cost) VALUES (?,?,?,?,?,?,?,?,?,?)""",(cid,now_iso(),f.get('technician','').strip(),f['description'].strip(),method,hours,rate,labor,parts,other))
   elif action=='stage':
    stage=f.get('stage')
    if stage not in ('Needs Repair','In Repair','QC'):raise ValueError('Invalid stage change')
    c.execute('UPDATE core_items SET stage=?,qc_result=NULL WHERE id=?',(stage,cid))
   elif action=='qc':
    result=f.get('result');tech=f.get('technician','').strip()
    if core['stage']!='QC':raise ValueError('Move the core to QC first')
    if result not in ('Passed','Failed') or not tech:raise ValueError('QC result and technician required')
    c.execute('UPDATE core_items SET qc_result=?,qc_technician=?,qc_date=?,qc_notes=?,stage=? WHERE id=?',(result,tech,now_iso(),f.get('notes','').strip(),'Ready for Stock' if result=='Passed' else 'In Repair',cid))
   elif action=='convert':
    if not can('core_convert'):raise ValueError('Only an Admin can convert a core into stock')
    if core['stage']!='Ready for Stock' or core['qc_result']!='Passed':raise ValueError('Passing QC is required before conversion')
    repair_cost=c.execute('SELECT COALESCE(SUM(labor_cost+parts_cost+other_cost),0) FROM core_transactions WHERE core_id=?',(cid,)).fetchone()[0]
    code=new_inventory_code(c,'CC');total=round(core['original_cost']+repair_cost,2)
    cur=c.execute("""INSERT INTO units(manufacturer,product_raw,product_key,serial,cost,notes,location_raw,physical_location,status,condition,review,inventory_code,supplier,received_date)
     VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(core['manufacturer'],core['product_raw'],norm(core['product_raw']),core['serial'],total,'Converted from '+core['inventory_code']+'; '+(core['notes'] or ''),core['location'],core['location'],'Active','Refurbished','',code,core['supplier'],date.today().isoformat()))
    c.execute('UPDATE core_items SET stage=?,converted_unit_id=? WHERE id=?',('Converted to Stock',cur.lastrowid,cid))
    flash('Converted to '+code+' with total cost $'+format(total,',.2f'))
   else:raise ValueError('Unknown action')
   c.commit()
  except (ValueError,sqlite3.Error) as e:c.rollback();c.close();return str(e),400
  c.close();return redirect(url_for('core_detail',cid=cid))
 core=c.execute('SELECT * FROM core_items WHERE id=?',(cid,)).fetchone()
 tx=c.execute('SELECT * FROM core_transactions WHERE core_id=? ORDER BY id DESC',(cid,)).fetchall()
 total=core['original_cost']+sum(t['labor_cost']+t['parts_cost']+t['other_cost'] for t in tx)
 c.close();return render_template_string(CORE_TEMPLATE,detail=core,transactions=tx,total=total,rows=[])

TEMPLATE=add_user_block((BASE/'dashboard.html').read_text(encoding='utf-8'))
ENTRY_TEMPLATE=add_user_block((BASE/'inventory_entry.html').read_text(encoding='utf-8'))
CORE_TEMPLATE=add_user_block((BASE/'cores.html').read_text(encoding='utf-8'))
@app.route('/')
def index():
 view=request.args.get('view','Active');q=request.args.get('q','').strip();page=max(1,request.args.get('page',1,type=int) or 1);c=conn();where=[];params=[]
 if view=='Active':where.append("status IN ('Active','Rental')")
 elif view!='All':where.append('status=?');params.append(view)
 if q:where.append('(product_raw LIKE ? OR serial LIKE ? OR notes LIKE ? OR location_raw LIKE ?)');params.extend(['%'+q+'%']*4)
 clause=' WHERE '+' AND '.join(where) if where else ''
 stats=c.execute('SELECT COUNT(*) n,COUNT(cost) priced,COALESCE(SUM(cost),0) val FROM units'+clause,params).fetchone();overview={r['status']:r['n'] for r in c.execute('SELECT status, COUNT(*) n FROM units GROUP BY status')};all_count=sum(overview.values());rows=c.execute('SELECT * FROM units'+clause+' ORDER BY id LIMIT 300 OFFSET ?',params+[(page-1)*300]).fetchall();c.close()
 return render_template_string(TEMPLATE,view=view,q=q,page=page,total=stats['n'],priced=stats['priced'],value=stats['val'],units=rows,unit=None,overview=overview,all_count=all_count)
@app.route('/unit/<int:uid>')
def unit(uid):
 c=conn();u=c.execute('SELECT * FROM units WHERE id=?',(uid,)).fetchone()
 if not u:c.close();return 'Not found',404
 moves=c.execute('SELECT * FROM movements WHERE unit_id=? ORDER BY id DESC',(uid,)).fetchall()
 unit_repairs=c.execute('SELECT * FROM repair_orders WHERE unit_id=? ORDER BY id DESC',(uid,)).fetchall()
 same=serial_matches(c,u['serial'],exclude=('unit',uid));c.close()
 return render_template_string(TEMPLATE,view=u['status'],unit=u,overview={},all_count=0,movements=moves,unit_repairs=unit_repairs,same_serial=same)
@app.post('/transfer/<int:uid>')
def transfer(uid):
 new=norm(request.form.get('location'));status,physical=classify(new)
 if status!='Active':flash('Transfer requires a recognized physical location');return redirect(url_for('unit',uid=uid))
 c=conn();u=c.execute('SELECT physical_location,status FROM units WHERE id=?',(uid,)).fetchone()
 if not u or u['status']=='Sold archive':c.close();return 'Sold records cannot be transferred',400
 c.execute('INSERT INTO movements(unit_id,old_location,new_location,note) VALUES (?,?,?,?)',(uid,u['physical_location'],physical,(request.form.get('note','').strip()+' ['+g.user['username']+']').strip()));c.execute('UPDATE units SET physical_location=?,status=? WHERE id=?',(physical,'Active',uid));c.commit();c.close();flash('Transfer recorded');return redirect(url_for('unit',uid=uid))
@app.post('/repair/<int:uid>')
def repair(uid):
 c=conn();u=c.execute('SELECT status FROM units WHERE id=?',(uid,)).fetchone()
 if not u or u['status']=='Sold archive':c.close();return 'Sold records cannot receive new repairs in this prototype',400
 c.execute('INSERT INTO repair_orders(unit_id,repair_type,notes) VALUES (?,?,?)',(uid,request.form.get('type'),(request.form.get('notes','').strip()+' ['+g.user['username']+']').strip()));c.commit();c.close();flash('Repair record created');return redirect(url_for('unit',uid=uid))
@app.get('/admin/backup')
def backup_download():
 # Admins can download a consistent copy of the whole database.
 if not can('backup'):abort(403)
 import tempfile
 tmp=Path(tempfile.mkdtemp())/('ccmedical-backup-'+datetime.now().strftime('%Y%m%d-%H%M')+'.db')
 src=conn();dst=sqlite3.connect(tmp);src.backup(dst);dst.close();src.close()
 return send_file(tmp,as_attachment=True,download_name=tmp.name)

RESTORE_PAGE='''<p><a href="/products/admin">Back to admin</a></p><div class="panel"><h1>Restore database from a backup</h1>
<p>Upload an <b>inventory.db</b> file (for example, the one from your computer, or a backup downloaded from this site). It replaces <b>all</b> current inventory data. A safety copy of the current database is downloaded automatically first, so keep it until you've checked the restored data.</p>
<p>User accounts on this site are kept if the uploaded file has none.</p>
<form method="post" enctype="multipart/form-data"><label>Database file</label><input type="file" name="file" accept=".db,.sqlite,.sqlite3" required>
<label><input type="checkbox" name="confirm" value="yes" required style="width:auto"> I understand this replaces all current data</label>
<button class="danger">Restore database</button></form></div>'''

@app.route('/admin/restore',methods=['GET','POST'])
def restore_upload():
 # Admin-only: replace the live database with an uploaded SQLite file.
 from auth import simple_page
 if not can('restore'):abort(403)
 if request.method=='GET':return simple_page('Restore database',RESTORE_PAGE,width=640)
 f=request.files.get('file')
 if not f or request.form.get('confirm')!='yes':flash('Choose a file and tick the confirmation box.');return redirect(url_for('restore_upload'))
 import tempfile
 work=Path(tempfile.mkdtemp());up=work/'upload.db';f.save(up)
 try:
  if up.read_bytes()[:16]!=b'SQLite format 3\x00':raise ValueError('That file is not a SQLite database.')
  t=sqlite3.connect(up)
  try:
   if t.execute('PRAGMA integrity_check').fetchone()[0]!='ok':raise ValueError('That database file is damaged.')
   tables={r[0] for r in t.execute("SELECT name FROM sqlite_master WHERE type='table'")}
   if 'units' not in tables:raise ValueError('That file is not a CC Medical inventory database.')
  finally:t.close()
 except (ValueError,sqlite3.DatabaseError) as e:
  flash(str(e) if isinstance(e,ValueError) else 'That file could not be read as a database.');return redirect(url_for('restore_upload'))
 # Keep a safety copy of the current data, then copy the upload into the live database.
 safety=work/('before-restore-'+datetime.now().strftime('%Y%m%d-%H%M')+'.db')
 live=conn();s_db=sqlite3.connect(safety);live.backup(s_db);s_db.close()
 current_users=[dict(r) for r in live.execute('SELECT * FROM users')]
 src=sqlite3.connect(up);src.backup(live);src.close();live.close()
 # If the uploaded file has no accounts, copy the current ones back BEFORE anything else creates users,
 # so everyone keeps the same logins (and the same account IDs).
 c=conn();ensure_user_schema(c)
 file_has_users=c.execute('SELECT COUNT(*) FROM users').fetchone()[0]>0
 if not file_has_users:
  have={r[1] for r in c.execute('PRAGMA table_info(users)')}
  for u in current_users:
   cols=[k for k in u.keys() if k in have];c.execute('INSERT INTO users('+','.join(cols)+') VALUES ('+','.join('?'*len(cols))+')',[u[k] for k in cols])
 c.commit();c.close()
 init_db()
 # The restored file brought its own accounts (IDs may differ), so sign in again.
 if file_has_users:session.clear()
 return send_file(safety,as_attachment=True,download_name=safety.name)

from product_database import register_product_routes
register_auth(app,conn)
register_product_routes(app,conn,BASE)
register_client_routes(app,conn)
R.register_repair_routes(app,conn,BASE,new_inventory_code)

def init_db():
 c=conn()
 c.executescript(SCHEMA);R.ensure_repair_schema(c);ensure_entry_schema(c);ensure_user_schema(c);ensure_client_schema(c)
 c.execute("UPDATE units SET status='Rental' WHERE UPPER(TRIM(location_raw))='RENTAL' AND status='On loan'")
 # On a host without a console, the first admin can come from environment variables.
 # This only runs while there are no users at all.
 u,pw=os.environ.get('CCM_INITIAL_ADMIN_USER'),os.environ.get('CCM_INITIAL_ADMIN_PASSWORD')
 if u and pw and c.execute('SELECT COUNT(*) FROM users').fetchone()[0]==0:
  try:create_user(c,u,pw,is_admin=True);print('Created initial admin user',u)
  except ValueError as e:print('Initial admin not created:',e)
 seed_users_from_env(c)
 c.commit();c.close()
init_db()

def cli_create_admin(username):
 import getpass
 pw=getpass.getpass('Password for '+username+' (at least 10 characters): ')
 if pw!=getpass.getpass('Repeat password: '):print('Passwords do not match.');return 1
 c=conn()
 try:create_user(c,username,pw,is_admin=True);c.commit();print('Admin user created:',username)
 except ValueError as e:print(e);return 1
 finally:c.close()
 return 0

if __name__=='__main__':
 cmd=sys.argv[1] if len(sys.argv)>1 else ''
 if cmd=='import':
  print('Imported records, duplicate serial groups:',import_file(sys.argv[2]));sys.exit()
 if cmd=='create-admin':
  if len(sys.argv)<3:print('Usage: python app.py create-admin USERNAME');sys.exit(1)
  sys.exit(cli_create_admin(sys.argv[2]))
 c=conn();n=c.execute('SELECT COUNT(*) FROM users').fetchone()[0];c.close()
 if n==0:print('No users yet. Create one first:  python app.py create-admin yourname');sys.exit(1)
 # Local use over plain http, so turn off HTTPS-only cookies.
 app.config['SESSION_COOKIE_SECURE']=False
 app.run(host='127.0.0.1',port=5000,debug=False)
